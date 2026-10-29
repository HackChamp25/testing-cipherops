"""Build docs/lab-architecture.pptx: editable slides for the real-DNS lab architecture.

Same visual language as scripts/make_architecture_pptx.py (library harness deck).

Usage: pip install python-pptx==1.0.2 && python scripts/make_lab_architecture_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "lab-architecture.pptx"

NAVY = "1F3A5F"
GREY_TEXT = "4A5568"
STYLES = {
    "container": ("F4F6FA", "8A99B3", NAVY),
    "orchestrator": ("1F3A5F", "1F3A5F", "FFFFFF"),
    "step": ("2F6FB0", "2F6FB0", "FFFFFF"),
    "lib": ("E8F1FB", "2F6FB0", NAVY),
    "live": ("FFF4E5", "E08A00", "7A4A00"),
    "input": ("EEF7EE", "3C8D40", "1E4D21"),
    "output": ("F3EEF9", "7B52AB", "3F2560"),
    "band": ("FFFFFF", "8A99B3", NAVY),
    "note": ("FFFFFF", "C9D2E0", GREY_TEXT),
    "dash": ("EAF4FF", "2F6FB0", NAVY),
}


def box(slide, x, y, w, h, title, lines=(), style="step", title_pt=13, body_pt=10,
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    fill, line, text = STYLES[style]
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.color.rgb = RGBColor.from_string(line)
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.06))
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = title
    r.font.bold = True
    r.font.size = Pt(title_pt)
    r.font.color.rgb = RGBColor.from_string(text)
    for line_text in lines:
        p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(body_pt)
        r.font.color.rgb = RGBColor.from_string(text)
    return s


def arrow(slide, x1, y1, x2, y2, color=NAVY, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def text(slide, x, y, w, h, value, pt=12, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def architecture_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.3, 0.12, 12.7, 0.4, "Cipherops: dual-path architecture (library + real DNS lab)", pt=24, bold=True)
    text(s, 0.3, 0.52, 12.7, 0.3,
         "PS-036 · Same PQC (SLH-DSA) with and without MTL · three sizes · pinned Verisign reference · IETF / national panel",
         pt=11, color=GREY_TEXT)

    # Left: library path (companion)
    box(s, 0.3, 0.95, 4.0, 5.35, "Path A — Library harness (main Cipherops)",
        ["Pinned Fedora image · MTL @7cef019 · liboqs 0.16", "",
         "cipherops.runner  →  mtlbench.c", "Sign / verify in-process",
         "dnswire + cache model", "report.html  (results/runs/…)", "",
         "Best for: precise crypto timings,", "signature sizes, UDP-fit %"],
        style="container", title_pt=12, body_pt=10, anchor=MSO_ANCHOR.TOP)
    box(s, 0.45, 3.55, 3.7, 0.85, "Crypto (unmodified)",
        ["Verisign MTL · liboqs SLH-DSA · OpenSSL classical"],
        style="lib", title_pt=11, body_pt=9)
    box(s, 0.45, 4.55, 3.7, 1.5, "Outputs",
        ["trials · responses · workload · cache", "report.md / .html · figures/",
         "scripts/run_benchmark.ps1"],
        style="output", title_pt=11, body_pt=9, anchor=MSO_ANCHOR.TOP)

    # Center: dashboard
    box(s, 4.5, 0.95, 4.3, 1.15, "Dashboard (Streamlit)",
        ["live_lab/app.py · five tiles", "Home · Setup · Live · Benchmark · Compare · Reset"],
        style="dash", title_pt=13, body_pt=10)
    arrow(s, 4.3, 2.3, 4.5, 1.5, color="8A99B3")
    arrow(s, 8.8, 1.5, 9.0, 2.3, color="8A99B3")

    box(s, 4.5, 2.25, 4.3, 4.05, "Path B — Real DNS lab (ecosystem)",
        ["Stack draft06 (panel default) or latest", "MTL-mode NSD · Unbound · ldns tools", "",
         "Hierarchy: root → TLD → leaf zones", "Stub client → Unbound → authorities",
         "pcap on Unbound iface · unbound-control", "",
         "Benchmark suite: sizes 50 / 500 / 5,000", "self-contained (images only; no Prepare)",
         "suite report.html  (like Path A)"],
        style="container", title_pt=12, body_pt=10, anchor=MSO_ANCHOR.TOP)

    # Right: lab internals
    box(s, 9.0, 0.95, 4.0, 5.35, "Lab building blocks",
        style="container", title_pt=12, body_pt=9, anchor=MSO_ANCHOR.TOP)
    box(s, 9.15, 1.5, 3.7, 0.7, "ecolab (Python)",
        ["zones · sign · topology · run"], style="step", title_pt=11, body_pt=9)
    box(s, 9.15, 2.35, 3.7, 0.7, "Docker images",
        ["cipherops-eco-nsd / unbound / tools"], style="step", title_pt=11, body_pt=9)
    box(s, 9.15, 3.2, 3.7, 0.95, "Live tile",
        ["Query · cold/warm · pcap sequence", "Technical demonstration"], style="live", title_pt=11, body_pt=9)
    box(s, 9.15, 4.3, 3.7, 1.7, "Outputs",
        ["ecosystem/results/suites/<id>/", "report.md · report.html · figures/",
         "signing_*.csv · per-size run.json", "pcap/ · summary.csv"],
        style="output", title_pt=11, body_pt=9, anchor=MSO_ANCHOR.TOP)

    box(s, 0.3, 6.45, 12.75, 0.4,
        "Compared on equivalent workloads:   ECDSA  ·  Ed25519  ·  RSA-2048  ·  SLH-DSA  ·  SLH-DSA + MTL",
        style="band", title_pt=11)
    box(s, 0.3, 6.95, 12.75, 0.4,
        "Measured (lab): signing · ldns-verify-zone · pcap bytes · Unbound cache     |     "
        "Estimated: sample-trial sign/verify     |     Analytical: draft section 12 sizes",
        style="band", title_pt=11)


def flow_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.3, 0.12, 12.7, 0.4, "Cipherops lab: end-to-end Benchmark suite flow", pt=24, bold=True)
    box(s, 0.35, 0.6, 12.6, 0.45,
        "Dashboard Benchmark  →  python -m live_lab.suite params.json  →  "
        "for each size: sign · sample trials · up · measure · down  →  suite_report",
        style="orchestrator", title_pt=11)

    steps = [
        ("1  Configure", "sidebar + suite.py",
         ["Sizes 50 / 500 / 5k", "Scenarios (mtl vs slh-dsa)", "Trials · lookups · lan"],
         ["params.json", "_bench-N.yaml"]),
        ("2  Shape zones", "ecolab.zones",
         ["1 TLD, ≤4 leaves", "Split N names", "Frozen unsigned zones"],
         ["work/<tag>/zones"]),
        ("3  Sign once", "ecolab.sign",
         ["ldns-keygen / signzone", "jobs=1 clean timings", "Per-zone wall clock"],
         ["manifest.json", "signing_full.csv"]),
        ("4  Sign trials", "signtrials",
         ["Sample leaf zones", "Repeat N times", "Estimate + 95% CI"],
         ["signing_trials_N.csv"]),
        ("5  Bring up", "topology + compose",
         ["NSD hierarchy", "Unbound validating", "Stub client"],
         ["compose.yml", "containers"]),
        ("6  Measure", "ecolab.run",
         ["Cold then warm", "Seeded 2000 lookups", "pcap + cache stats"],
         ["summary.csv", "pcap/", "run.json"]),
        ("7  Report", "suite_report.py",
         ["Figures", "PS-036 tables", "MTL vs SLH ratios"],
         ["report.md / .html", "figures/", "→ Compare"]),
    ]
    w, gap, x0 = 1.7, 0.12, 0.35
    for i, (title, file, does, makes) in enumerate(steps):
        x = x0 + i * (w + gap)
        box(s, x, 1.25, w, 0.8, title, [file], style="step", title_pt=12, body_pt=9)
        box(s, x, 2.2, w, 1.7, "What it does", ["• " + d for d in does], style="note", title_pt=10,
            body_pt=10, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
        box(s, x, 4.05, w, 1.35, "Produces", makes, style="output", title_pt=10, body_pt=10,
            anchor=MSO_ANCHOR.TOP)
        if i:
            arrow(s, x - gap, 1.65, x, 1.65)

    box(s, 0.35, 5.6, 6.15, 1.55, "Companion: Path A (library)",
        ["scripts/run_benchmark.ps1  →  results/runs/<id>/report.html",
         "Use ratios (MTL / SLH-DSA) next to the lab report — not absolute seconds",
         "Library pin MTL 7cef019 · Lab default stack draft06 (condensed-by-default)"],
        style="lib", title_pt=12, body_pt=10, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
    box(s, 6.7, 5.6, 6.25, 1.55, "Demo & Compare",
        ["Live tile: one query + pcap sequence (technical demonstration)",
         "Compare tile: Measured / Estimated / Analytical badges · Across sizes",
         "Reset tile: wipe suites only when you choose — refresh does not delete"],
        style="live", title_pt=12, body_pt=10, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)


def dual_story_slide(prs):
    """Third slide: what to show the panel (outputs map)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.3, 0.12, 12.7, 0.4, "Cipherops: what the panel sees (outputs map)", pt=24, bold=True)
    text(s, 0.3, 0.52, 12.7, 0.28,
         "Keep core PS-036 numbers; use ratios to bridge library and lab; demo proves real DNS.",
         pt=11, color=GREY_TEXT)

    cols = [
        (0.3, "KEEP — core scorecard", "input",
         ["Signing time (3 sizes)", "Verification (labelled clock)", "Transmitted bytes",
          "Cache measured", "Analytical (separate)", "Pinned env + trials", "report.html both paths"]),
        (4.55, "KEEP — panel aha", "step",
         ["MTL vs SLH-DSA ratios", "UDP fit / truncation story", "Lab pcap wire bytes",
          "Across-sizes charts", "Live pcap 60s demo", "Measured-once vs Estimated"]),
        (8.8, "SOFT / bury", "note",
         ["Per-tier hop charts", "Wan / lossy / Zipf deep dive", "All five algos on one slide",
          "Jobs / sample knobs", "Absolute lab vs library seconds", "mixed-mtl unless asked"]),
    ]
    for x, title, style, lines in cols:
        box(s, x, 1.0, 4.05, 4.5, title, [""] + lines, style=style, title_pt=14, body_pt=12,
            anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)

    box(s, 0.3, 5.7, 12.75, 1.45, "One sentence for the room",
        ["Library proves the crypto cost.  Lab proves the DNS path.  "
         "Same algorithm with and without MTL; ratios agree; reports and pins are reproducible."],
        style="orchestrator", title_pt=14, body_pt=13)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    architecture_slide(prs)
    flow_slide(prs)
    dual_story_slide(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
